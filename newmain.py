import os
import json
import groq
import streamlit as st
import tempfile
import dateparser
import subprocess
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
from moviepy.video.io.VideoFileClip import VideoFileClip
import contextlib

# Directly set the environment variables
os.environ["GROQ_API_KEY"] = "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"
os.environ["ASSEMBLYAI_API_KEY"] = "c5f6175e88f64dcb8a9fdaef9bd413e1"

# Get API key from environment
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")

if not GROQ_API_KEY:
    raise ValueError("GROQ_API_KEY is missing.")

# Initialize Groq Client
client = groq.Client(api_key=GROQ_API_KEY)

# Available languages for summary translation
LANGUAGES = {
    "English": "English",
    "Spanish": "Spanish (Español)",
    "French": "French (Français)",
    "German": "German (Deutsch)",
    "Chinese": "Chinese (中文)",
    "Japanese": "Japanese (日本語)",
    "Russian": "Russian (Русский)",
    "Arabic": "Arabic (العربية)",
    "Hindi": "Hindi (हिन्दी)",
    "Portuguese": "Portuguese (Português)",
    "Italian": "Italian (Italiano)",
    "Dutch": "Dutch (Nederlands)",
    "Korean": "Korean (한국어)",
    "Swedish": "Swedish (Svenska)",
    "Turkish": "Turkish (Türkçe)"
}

def transcribe_with_assemblyai(audio_path):
    """Transcribe audio using AssemblyAI API."""
    try:
        import assemblyai as aai
        
        # Get AssemblyAI API key from environment or prompt user
        api_key = os.getenv("ASSEMBLYAI_API_KEY")
        if not api_key:
            api_key = st.text_input("Enter your AssemblyAI API key (free tier available at assemblyai.com)", type="password")
            if not api_key:
                return "Transcription failed: AssemblyAI API key required."
        
        # Initialize client
        aai.settings.api_key = api_key
        
        # Create transcriber
        transcriber = aai.Transcriber()
        
        with st.spinner("Uploading and transcribing audio (this may take some time)..."):
            # Start transcription
            transcript = transcriber.transcribe(audio_path)
            
            if transcript.status == "completed":
                return transcript.text
            else:
                return f"Transcription failed: {transcript.status}"
    except ImportError:
        st.error("AssemblyAI package not installed. Please install with: pip install assemblyai")
        return "Transcription failed: Missing dependency."
    except Exception as e:
        st.error(f"Transcription error: {str(e)}")
        return f"Transcription failed: {str(e)}"

def extract_text_from_file(uploaded_file):
    """Extract text from TXT, PDF, DOCX, audio, or video files."""
    if uploaded_file is None:
        return None
        
    if uploaded_file.type == "text/plain":
        return uploaded_file.read().decode("utf-8")
    elif uploaded_file.type == "application/pdf":
        pdf_reader = PdfReader(uploaded_file)
        return "\n".join([page.extract_text() for page in pdf_reader.pages if page.extract_text()])
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    elif uploaded_file.type.startswith("audio/"):
        # Handle audio files
        with st.spinner("Processing audio file..."):
            with tempfile.NamedTemporaryFile(delete=False, suffix=f".{uploaded_file.name.split('.')[-1]}") as tmp_file:
                tmp_file.write(uploaded_file.getvalue())
                audio_path = tmp_file.name
            
            transcript = transcribe_with_assemblyai(audio_path)
            
            # Clean up temp file
            try:
                os.unlink(audio_path)
            except:
                pass
                
            return transcript
    elif uploaded_file.type.startswith("video/"):
        # Handle video files
        with st.spinner("Extracting audio from video..."):
            try:
                # Save video to temporary file
                with tempfile.NamedTemporaryFile(delete=False, suffix=f".{uploaded_file.name.split('.')[-1]}") as tmp_video:
                    tmp_video.write(uploaded_file.getvalue())
                    video_path = tmp_video.name
                
                # Extract audio from video
                audio_path = tempfile.NamedTemporaryFile(delete=False, suffix=".wav").name
                try:
                    video = VideoFileClip(video_path)
                    video.audio.write_audiofile(audio_path, logger=None)
                    video.close()
                except Exception as e:
                    st.error(f"Error extracting audio: {e}")
                    return f"Failed to extract audio from video: {str(e)}"
                
                transcript = transcribe_with_assemblyai(audio_path)
                
                # Clean up temp files
                try:
                    os.unlink(video_path)
                    os.unlink(audio_path)
                except:
                    pass
                    
                return transcript
            except ImportError:
                st.error("MoviePy not installed. Please install with: pip install moviepy")
                return "Failed to process video: Missing moviepy dependency."
            except Exception as e:
                st.error(f"Error processing video: {str(e)}")
                return f"Video processing failed: {str(e)}"
    return None

def summarize_transcript(transcript, language="English"):
    """Use Groq API to summarize the meeting transcript in the specified language."""
    prompt = f"""
    Summarize the following meeting transcript in a concise and informative way.
    Highlight the key discussion points, decisions made, and important takeaways.
    
    IMPORTANT: Your summary must be written in {language} language.
    
    Transcript:
    {transcript}
    """
    try:
        response = client.chat.completions.create(
            model="llama3-70b-8192",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        st.error(f"Error summarizing transcript: {e}")
        return "Summarization failed."

def extract_action_items(transcript, language="English"):
    """Use Groq API to extract action items, responsible persons, and deadlines in the specified language."""
    prompt = f"""
    You are an AI that extracts structured action items from meeting transcripts.
    Identify action items, responsible persons, and deadlines where available.
    
    IMPORTANT: Translate the following information into {language} language:
    1. The 'action' description for each item
    2. Any other text that will be shown to the user
    
    Return the output as a JSON list with these keys: 'person', 'action', and 'deadline'.
    Keep person names in their original form, but translate the actions and explanatory text.

    Ensure the deadline is extracted accurately. If no deadline is mentioned, leave it empty.

    Transcript:
    {transcript}
    """
    try:
        response = client.chat.completions.create(
            model="llama3-70b-8192",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3
        )
        # Handle potential JSON parsing issues
        try:
            content = response.choices[0].message.content
            # Clean up the response to ensure it's valid JSON
            # Find JSON content between possible markdown code blocks
            import re
            json_match = re.search(r'```(?:json)?\s*([\s\S]*?)\s*```', content)
            if json_match:
                content = json_match.group(1)
            else:
                # Try to find array brackets directly
                json_match = re.search(r'\[\s*{[\s\S]*}\s*\]', content)
                if json_match:
                    content = json_match.group(0)
                    
            action_items = json.loads(content)
            
            # Handle non-list responses
            if not isinstance(action_items, list):
                if isinstance(action_items, dict) and 'items' in action_items:
                    action_items = action_items['items']
                else:
                    action_items = [action_items]

        except json.JSONDecodeError:
            st.warning("Could not parse JSON response. Displaying raw output.")
            st.code(response.choices[0].message.content)
            action_items = []

        # Parse dates using dateparser with language context
        for item in action_items:
            if item.get("deadline"):
                parsed_date = dateparser.parse(item["deadline"], languages=[language.lower()])
                item["deadline"] = parsed_date.strftime("%Y-%m-%d") if parsed_date else item["deadline"]
        return action_items
    except Exception as e:
        st.error(f"Error extracting action items: {e}")
        return []
    
def generate_pdf(summary, action_items, language="English"):
    """Generate a PDF report for download in the specified language."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    # Set font that supports multiple languages
    try:
        # Try to use a font that supports the language
        if language in ["Chinese", "Japanese", "Korean"]:
            pdf.add_font('NotoSansCJK', '', 'NotoSansCJK-Regular.ttf', uni=True)
            pdf.set_font('NotoSansCJK', '', 14)
        elif language in ["Arabic", "Hindi"]:
            pdf.add_font('NotoSansArabic', '', 'NotoSansArabic-Regular.ttf', uni=True)
            pdf.set_font('NotoSansArabic', '', 14)
        else:
            pdf.set_font("Arial", style="B", size=16)
    except:
        # Fallback to Arial which has decent support for Western languages
        pdf.set_font("Arial", style="B", size=16)
    
    # Get title translations
    titles = {
        "English": {"report": "Meeting Summary Report", "summary": "Summary", "actions": "Key Action Items"},
        "Spanish": {"report": "Informe de Resumen de Reunión", "summary": "Resumen", "actions": "Elementos de Acción Clave"},
        "French": {"report": "Rapport de Synthèse de Réunion", "summary": "Résumé", "actions": "Points d'Action Clés"},
        "German": {"report": "Besprechungszusammenfassung", "summary": "Zusammenfassung", "actions": "Wichtige Maßnahmen"},
        "Chinese": {"report": "会议摘要报告", "summary": "摘要", "actions": "关键行动项目"},
        "Japanese": {"report": "会議要約レポート", "summary": "要約", "actions": "主要なアクションアイテム"},
        "Russian": {"report": "Отчет о Резюме Встречи", "summary": "Резюме", "actions": "Ключевые Пункты Действий"},
        "Arabic": {"report": "تقرير ملخص الاجتماع", "summary": "ملخص", "actions": "بنود العمل الرئيسية"},
        "Hindi": {"report": "बैठक सारांश रिपोर्ट", "summary": "सारांश", "actions": "प्रमुख कार्य आइटम"},
        "Portuguese": {"report": "Relatório de Resumo da Reunião", "summary": "Resumo", "actions": "Itens de Ação Chave"},
        "Italian": {"report": "Rapporto di Sintesi della Riunione", "summary": "Sintesi", "actions": "Punti di Azione Chiave"},
        "Dutch": {"report": "Vergadering Samenvatting Rapport", "summary": "Samenvatting", "actions": "Belangrijke Actiepunten"},
        "Korean": {"report": "회의 요약 보고서", "summary": "요약", "actions": "주요 작업 항목"},
        "Swedish": {"report": "Mötessammanfattningsrapport", "summary": "Sammanfattning", "actions": "Viktiga Åtgärdspunkter"},
        "Turkish": {"report": "Toplantı Özeti Raporu", "summary": "Özet", "actions": "Önemli Eylem Maddeleri"}
    }
    
    # Default to English if language not in dictionary
    lang_titles = titles.get(language, titles["English"])
    
    pdf.cell(200, 10, lang_titles["report"], ln=True, align="C")
    pdf.set_font("Arial", size=12)
    pdf.ln(10)
    pdf.cell(200, 10, f"{lang_titles['summary']}:", ln=True)
    pdf.multi_cell(0, 10, summary)
    pdf.ln(10)
    pdf.cell(200, 10, f"{lang_titles['actions']}:", ln=True)
    for item in action_items:
        pdf.multi_cell(0, 10, f"- {item['person']}: {item['action']} (Deadline: {item['deadline']})")
    pdf_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(pdf_file.name)
    return pdf_file.name

# Streamlit UI
st.set_page_config(page_title="AI Meeting Summarizer", layout="wide", page_icon="📝")
st.markdown("<h1 style='text-align: center;'>📋 AI Meeting Summarizer</h1>", unsafe_allow_html=True)
st.markdown("<h4 style='text-align: center;'>Upload a meeting transcript or recording</h4>", unsafe_allow_html=True)
st.markdown("---")

# File upload area with multiple file types
uploaded_file = st.file_uploader("Upload a transcript or recording", 
                                type=["txt", "pdf", "docx", "mp3", "wav", "m4a", "mp4", "mov", "avi"])

# Add language selection
selected_language = st.selectbox(
    "Select summary language",
    options=list(LANGUAGES.keys()),
    format_func=lambda x: LANGUAGES[x],
    index=0
)

if uploaded_file:
    # Process and display transcript
    transcript = extract_text_from_file(uploaded_file)
    
    if transcript:
        st.subheader("📝 Transcript Preview")
        st.text_area("Transcript", transcript, height=200)
        
        # Allow user to edit the transcript before processing
        edited_transcript = st.text_area("Edit Transcript (if needed)", transcript, height=300)
        
        # Generate summary and action items
        if st.button("🚀 Generate Summary & Action Items"):
            with st.spinner(f"Analyzing transcript and generating {selected_language} summary..."):
                summary = summarize_transcript(edited_transcript, selected_language)
                action_items = extract_action_items(edited_transcript, selected_language)
            
            st.subheader("📌 Summary")
            st.write(summary)
            
            st.subheader("✅ Key Action Items")
            if action_items:
                st.table(action_items)
            else:
                st.write("No clear action items detected.")
            
            # Generate and offer PDF download
            pdf_path = generate_pdf(summary, action_items, selected_language)
            with open(pdf_path, "rb") as pdf_file:
                st.download_button(label="📥 Download Summary (PDF)", 
                                  data=pdf_file, 
                                  file_name=f"Meeting_Summary_{selected_language}.pdf", 
                                  mime="application/pdf")
    else:
        st.error("Failed to extract text from the uploaded file.")
else:
    st.info("Please upload a file to begin.")